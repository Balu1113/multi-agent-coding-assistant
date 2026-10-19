from __future__ import annotations

import csv
import io
import logging
import os
import re
from pathlib import Path

logger = logging.getLogger(__name__)

_OFFICE_MIME_TYPES: dict[str, str] = {
    "application/msword": "word",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "word",
    "application/vnd.ms-excel": "excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "excel",
    "application/vnd.ms-powerpoint": "powerpoint",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "powerpoint",
    "text/csv": "csv",
    "application/pdf": "pdf",
}


def classify_mime(mime: str) -> str:
    mime_lower = mime.lower().strip()
    if mime_lower.startswith("image/"):
        return "image"
    if mime_lower.startswith("video/"):
        return "video"
    if mime_lower.startswith("audio/"):
        return "audio"
    if mime_lower == "application/pdf":
        return "pdf"
    if mime_lower.startswith("text/"):
        return "text"
    return _OFFICE_MIME_TYPES.get(mime_lower, "other")


def extract_text(file_path: str | os.PathLike, mime: str) -> str:
    path = Path(file_path)
    if not path.exists():
        return ""
    cls = classify_mime(mime)
    try:
        if cls == "pdf":
            return _extract_pdf(path)
        if cls == "word":
            return _extract_docx(path)
        if cls == "excel":
            return _extract_xlsx(path)
        if cls == "csv":
            return _extract_csv(path)
        if cls == "powerpoint":
            return _extract_pptx(path)
        if cls == "text":
            return _extract_text(path)
        if cls in ("image", "video", "audio", "other"):
            return ""
        return _extract_text(path)
    except Exception as e:
        logger.warning("Failed to extract text from %s: %s", path.name, e)
        return ""


def _extract_pdf(path: Path) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text.strip())
        return "\n\n".join(pages)
    except Exception as e:
        logger.warning("PyPDF2 failed: %s", e)
        return ""


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paras)
    except Exception as e:
        logger.warning("python-docx failed: %s", e)
        return ""


def _extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines.append(f"=== Sheet: {sheet_name} ===")
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(c) if c is not None else "" for c in row]
                lines.append("\t".join(row_vals))
        return "\n".join(lines)
    except Exception as e:
        logger.warning("openpyxl failed: %s", e)
        return ""


def _extract_csv(path: Path) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            lines = ["\t".join(row) for row in reader]
        return "\n".join(lines)
    except Exception as e:
        logger.warning("CSV extraction failed: %s", e)
        return ""


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        lines = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        slide_texts.append(" | ".join(cells))
            if slide_texts:
                lines.append(f"--- Slide {slide_num} ---")
                lines.extend(slide_texts)
        return "\n".join(lines)
    except Exception as e:
        logger.warning("python-pptx failed: %s", e)
        return ""


def _extract_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        try:
            return path.read_text(encoding="latin-1")
        except Exception as e:
            logger.warning("Text extraction failed: %s", e)
            return ""
